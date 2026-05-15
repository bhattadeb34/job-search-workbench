"""AI helpers — uses Gemini REST API directly (no google-generativeai package needed)."""
import io
import os
from typing import List

import pandas as pd
import requests
from bs4 import BeautifulSoup

_GEMINI_URL = (
    "https://generativelanguage.googleapis.com/v1beta"
    "/models/gemini-2.5-flash:generateContent"
)


def _generate(prompt: str) -> str:
    key = os.environ.get("GEMINI_API_KEY", "")
    if not key:
        raise RuntimeError("GEMINI_API_KEY not configured.")
    r = requests.post(
        f"{_GEMINI_URL}?key={key}",
        json={"contents": [{"parts": [{"text": prompt}]}]},
        timeout=90,
    )
    r.raise_for_status()
    return r.json()["candidates"][0]["content"]["parts"][0]["text"]


# ── Profile extraction ────────────────────────────────────────────────────────

def extract_text_from_upload(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    if name.endswith(".txt") or name.endswith(".md"):
        return data.decode("utf-8", errors="replace")
    if name.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"Unsupported file type: {uploaded_file.name}")


def fetch_url_text(url: str) -> str:
    r = requests.get(url, timeout=15, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)[:6000]


# ── Keyword suggestions ───────────────────────────────────────────────────────

def suggest_keywords(profile_text: str, existing: List[str]) -> List[str]:
    existing_str = "\n".join(existing[:20]) if existing else "(none yet)"
    prompt = (
        "You are a job search expert helping this candidate find relevant jobs. "
        "suggest 15-20 specific job search keyword phrases.\n\n"
        f"Candidate profile:\n{profile_text[:4000]}\n\n"
        f"Already using:\n{existing_str}\n\n"
        "Focus on: job titles, tools/software, methodologies, and industries that match their background. "
        "Return ONLY a plain list, one phrase per line, no numbering, no explanations."
    )
    resp = _generate(prompt)
    return [l.strip() for l in resp.strip().splitlines() if l.strip()]


# ── Job scoring ───────────────────────────────────────────────────────────────

def score_jobs(df: pd.DataFrame, profile_text: str, max_jobs: int = 40) -> pd.DataFrame:
    work = df.head(max_jobs).copy().reset_index(drop=True)
    BATCH = 10
    scores, reasons = [], []

    for i in range(0, len(work), BATCH):
        batch = work.iloc[i: i + BATCH]
        jobs_block = ""
        for j, (_, row) in enumerate(batch.iterrows()):
            desc = str(row.get("description", ""))[:300]
            jobs_block += f"\nJOB {j+1}: {row.get('title','?')} at {row.get('institution','?')}\n{desc}\n"

        prompt = (
            f"Rate each job for this candidate (1–10) and give a one-sentence reason.\n\n"
            f"Candidate profile:\n{profile_text[:1500]}\n\n"
            f"Jobs:\n{jobs_block}\n"
            f"Reply with exactly {len(batch)} lines:\n"
            "JOB 1: SCORE=X | REASON=...\nJOB 2: SCORE=X | REASON=...\netc."
        )
        try:
            resp = _generate(prompt)
            parsed = []
            for line in resp.strip().splitlines():
                if "SCORE=" in line and "REASON=" in line:
                    try:
                        s = int(line.split("SCORE=")[1].split("|")[0].strip())
                    except Exception:
                        s = 5
                    r = line.split("REASON=")[1].strip()
                    parsed.append((s, r))
            while len(parsed) < len(batch):
                parsed.append((5, "Could not parse"))
            for s, r in parsed[: len(batch)]:
                scores.append(s)
                reasons.append(r)
        except Exception:
            scores.extend([5] * len(batch))
            reasons.extend(["Scoring error"] * len(batch))

    result = work.copy()
    result["ai_score"] = scores
    result["ai_reason"] = reasons
    return result.sort_values("ai_score", ascending=False).reset_index(drop=True)


# ── Keyword quality review ────────────────────────────────────────────────────

def review_keywords(keywords: List[str], profile_text: str = "") -> dict:
    """Return structured keyword quality feedback."""
    import json as _json
    kw_block = "\n".join(f"- {k}" for k in keywords)
    profile_block = f"\nCandidate profile (for context):\n{profile_text[:1500]}" if profile_text else ""
    prompt = (
        "You are a job search expert. Review this keyword list and return a JSON object.\n\n"
        f"Keywords:\n{kw_block}{profile_block}\n\n"
        "Return ONLY valid JSON, no markdown, no explanation:\n"
        "{\n"
        '  "quality_score": <1-10>,\n'
        '  "summary": "<2-sentence overall assessment>",\n'
        '  "weak": ["<vague or unlikely to return good results>"],\n'
        '  "redundant": [["<kw1>", "<kw2 that means the same>"]],\n'
        '  "suggestions": ["<5-10 strong keywords to add based on the profile>"]\n'
        "}"
    )
    raw = _generate(prompt)
    # Strip markdown fences if present
    raw = raw.strip().strip("```json").strip("```").strip()
    try:
        return _json.loads(raw)
    except Exception:
        # Fallback: return minimal structure
        return {"quality_score": None, "summary": raw[:300], "weak": [], "redundant": [], "suggestions": []}


# ── Cover letter ──────────────────────────────────────────────────────────────

def draft_cover_letter(job: dict, profile_text: str) -> str:
    prompt = (
        "Write a concise, professional 3-paragraph cover letter.\n\n"
        f"Candidate profile:\n{profile_text[:2000]}\n\n"
        f"Position: {job.get('title','N/A')} at {job.get('institution','N/A')}\n"
        f"Description:\n{str(job.get('description',''))[:800]}\n\n"
        "Be specific. No generic filler. No placeholders like [Your Name]."
    )
    return _generate(prompt)


# ── Fit analysis ──────────────────────────────────────────────────────────────

def explain_fit(job: dict, profile_text: str) -> str:
    prompt = (
        "Analyze this job fit.\n\n"
        f"Candidate:\n{profile_text[:1500]}\n\n"
        f"Job: {job.get('title','N/A')} at {job.get('institution','N/A')}\n"
        f"Description:\n{str(job.get('description',''))[:800]}\n\n"
        "Give:\n- 3 bullets: GOOD fit reasons\n- 3 bullets: GAPS or concerns\n"
        "- Verdict: **Apply** / **Maybe** / **Skip** with one-line reason."
    )
    return _generate(prompt)
